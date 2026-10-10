from __future__ import annotations
"""
fundamentals.py - reads screener.in "Export to Excel" files (../screener_xlsx/<SYMBOL>.xlsx,
made by your screener_downloader.py) and returns the raw inputs for Steps 2b and 2c.

All money values are Rs crore, as in the Screener export.
Live use only: Screener restates history, so the backtest must use BSE XBRL
filings as first filed (doc fix #14).

COGS check (Screener's own formula, verified on AARTIIND FY26):
    COGS = Raw Material + Power & Fuel + Other Mfr. Exp - Change in Inventory
"""
import os
import time
from functools import lru_cache

import numpy as np
import openpyxl
import pandas as pd

import config as C

SECTIONS = {"PROFIT & LOSS": "pl", "Quarters": "q", "BALANCE SHEET": "bs", "CASH FLOW:": "cf",
            "PRICE:": "price", "DERIVED:": "dv", "META": "meta"}

NEG_INF, POS_INF = -np.inf, np.inf     # "rank last" / "rank best" markers for percentile ranking


def _read_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Data Sheet" not in wb.sheetnames:
        return None
    ws = wb["Data Sheet"]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _parse(path):
    rows = _read_sheet(path)
    if not rows:
        return None
    sec, D, meta = None, {}, {}
    for row in rows:
        lab = row[0] if row else None
        if isinstance(lab, str) and lab.strip() in SECTIONS:
            sec = SECTIONS[lab.strip()]
            D.setdefault(sec, {"dates": None, "rows": {}})
            if sec == "price":
                D[sec]["rows"]["price"] = row[1:]
            continue
        if sec is None or lab is None:
            continue
        L = str(lab).strip()
        if sec == "meta":
            meta[L] = row[1] if len(row) > 1 else None
        elif L == "Report Date":
            D[sec]["dates"] = row[1:]
        else:
            D[sec]["rows"].setdefault(L, row[1:])      # first 'Total' = total liabilities

    def frame(key, names):
        if key not in D or not D[key]["dates"]:
            return pd.DataFrame()
        pos = [i for i, d in enumerate(D[key]["dates"]) if d is not None]   # companies with < 10 years
        dates = [D[key]["dates"][i] for i in pos]                         # have blank leading columns
        out = {}
        for src in names:
            row = list(D[key]["rows"].get(src, []))
            v = [row[i] if i < len(row) else None for i in pos]
            out[src] = pd.to_numeric(pd.Series(v, dtype=object), errors="coerce").values
        return pd.DataFrame(out, index=pd.to_datetime(dates))

    pl = frame("pl", ["Sales", "Raw Material Cost", "Change in Inventory", "Power and Fuel", "Other Mfr. Exp",
                      "Other Income", "Depreciation", "Interest", "Profit before tax", "Net profit",
                      "Dividend Amount"])
    bs = frame("bs", ["Equity Share Capital", "Reserves", "Borrowings", "Other Liabilities", "Total",
                      "Net Block", "Capital Work in Progress", "Investments", "Other Assets", "Cash & Bank",
                      "No. of Equity Shares"])
    cf = frame("cf", ["Cash from Operating Activity"])
    dv = None
    if "dv" in D and D["dv"]["rows"].get("Adjusted Equity Shares in Cr") is not None and D.get("pl", {}).get("dates"):
        pos = [i for i, d in enumerate(D["pl"]["dates"]) if d is not None]
        row = list(D["dv"]["rows"]["Adjusted Equity Shares in Cr"])
        dv = pd.to_numeric(pd.Series([row[i] if i < len(row) else None for i in pos], dtype=object),
                           errors="coerce")
        dv.index = pl.index
    q = frame("q", ["Sales", "Net profit", "Depreciation", "Operating Profit"])
    return {"pl": pl, "bs": bs, "cf": cf, "adj_shares": dv, "q": q, "meta": meta}


def _ttm(q: pd.DataFrame, col: str):
    if q is None or q.empty or col not in q:
        return np.nan
    s = q[col].dropna()
    if len(s) < 4:
        return np.nan
    return float(s.iloc[-4:].sum())


def _avg_last(s: pd.Series, n: int):
    s = s.replace([np.inf, -np.inf], np.nan).dropna()
    if len(s) < max(1, n - 1):          # allow one missing year
        return np.nan
    return float(s.iloc[-n:].mean())


def _div(a, b):
    return np.nan if (b is None or not np.isfinite(b) or b == 0 or not np.isfinite(a)) else a / b


@lru_cache(maxsize=None)
def load(symbol: str) -> dict | None:
    path = os.path.join(C.SCREENER_XLSX_DIR, f"{symbol}.xlsx")
    if not os.path.exists(path):
        return None
    try:
        P = _parse(path)
    except Exception:
        return None
    if P is None or P["pl"].empty or P["bs"].empty:
        return None
    pl, bs, cf, q = P["pl"], P["bs"], P["cf"], P["q"]
    A = pl.join(bs, how="outer").join(cf, how="outer").sort_index()
    A = A[A["Sales"].notna() | A["Net profit"].notna()]
    if A.empty or (pd.Timestamp.now() - A.index[-1]).days > 550:
        return None                                        # empty or stale export (> ~18 months)
    g = lambda c: A[c] if c in A else pd.Series(np.nan, index=A.index)

    equity = g("Equity Share Capital").fillna(0) + g("Reserves").fillna(0)
    equity = equity.where(equity != 0)
    total = g("Total")
    borrow = g("Borrowings").fillna(0)
    cash = g("Cash & Bank").fillna(0)
    np_ = g("Net profit")
    pbt, intr, dep = g("Profit before tax"), g("Interest").fillna(0), g("Depreciation").fillna(0)
    oth_inc = g("Other Income").fillna(0)
    sales = g("Sales")
    cogs = (g("Raw Material Cost").fillna(0) + g("Power and Fuel").fillna(0)
            + g("Other Mfr. Exp").fillna(0) - g("Change in Inventory").fillna(0))
    cfo = g("Cash from Operating Activity")
    divd = g("Dividend Amount").fillna(0)
    ebit = pbt + intr
    ebitda = pbt + intr + dep
    op = ebitda - oth_inc                                  # operating EBITDA
    roce = ebit / (equity + borrow).where((equity + borrow) > 0)
    roe = np_ / equity.where(equity > 0)
    roa = np_ / total.where(total > 0)

    last = A.index[-1]
    L = lambda s: float(s.loc[last]) if (last in s.index and pd.notna(s.loc[last])) else np.nan
    pat3, cfo3, div3 = np_.iloc[-3:].sum(min_count=2), cfo.iloc[-3:].sum(min_count=2), divd.iloc[-3:].sum()

    # ---- quality inputs, non-financials (doc 2b, with edge-case rules)
    q1 = _div(L(sales) - L(cogs), L(total))
    q2 = _avg_last(roce, 3)
    q3 = NEG_INF if not (pat3 > 0) else min(C.Q3_CAP, _div(cfo3, pat3))
    q4 = NEG_INF if not (pat3 > 0) else min(C.Q4_CAP, _div(div3, pat3))
    nd, eb = L(borrow) - L(cash), L(ebitda)
    if np.isfinite(nd) and nd <= 0:
        q5 = POS_INF                                       # net cash -> best
    elif not (eb > 0):
        q5 = NEG_INF                                       # debt with no EBITDA -> last
    else:
        q5 = -nd / eb

    # ---- financials
    roe3, roa3 = _avg_last(roe, 3), _avg_last(roa, 3)
    payout = NEG_INF if not (pat3 > 0) else min(C.Q4_CAP, _div(div3, pat3))
    roe_sd5 = float(roe.dropna().iloc[-5:].std()) if roe.dropna().size >= 3 else np.nan

    # ---- Piotroski F-score (non-financials), signals computable from the export
    fs, fn = 0, 0
    if len(A) >= 2:
        p = A.index[-2]
        P_ = lambda s: float(s.loc[p]) if pd.notna(s.loc[p]) else np.nan
        tests = [
            (L(roa), lambda: L(roa) > 0),
            (L(cfo), lambda: L(cfo) > 0),
            (L(roa) - P_(roa), lambda: L(roa) > P_(roa)),
            (L(cfo) - L(np_), lambda: L(cfo) > L(np_)),
            (_div(L(borrow), L(total)) - _div(P_(borrow), P_(total)),
             lambda: _div(L(borrow), L(total)) <= _div(P_(borrow), P_(total))),
            (_div(L(g("Other Assets")), L(g("Other Liabilities"))) - _div(P_(g("Other Assets")), P_(g("Other Liabilities"))),
             lambda: _div(L(g("Other Assets")), L(g("Other Liabilities"))) > _div(P_(g("Other Assets")), P_(g("Other Liabilities")))),
            (_div(L(sales) - L(cogs), L(sales)) - _div(P_(sales) - P_(cogs), P_(sales)),
             lambda: _div(L(sales) - L(cogs), L(sales)) > _div(P_(sales) - P_(cogs), P_(sales))),
            (_div(L(sales), L(total)) - _div(P_(sales), P_(total)),
             lambda: _div(L(sales), L(total)) > _div(P_(sales), P_(total))),
        ]
        for val, ok in tests:
            if np.isfinite(val):
                fn += 1
                fs += int(ok())
        sh = P["adj_shares"]
        if sh is not None and sh.dropna().size >= 2:
            fn += 1
            fs += int(sh.dropna().iloc[-1] <= sh.dropna().iloc[-2] * 1.01)

    # ---- valuation raw inputs (mcap is added later from the live price)
    capex = np.nan
    if len(A) >= 2:
        nb = g("Net Block").fillna(0) + g("Capital Work in Progress").fillna(0)
        capex = float(nb.iloc[-1] - nb.iloc[-2] + dep.iloc[-1])
    shares = P["meta"].get("Number of shares")
    shares = float(shares) if isinstance(shares, (int, float)) and shares > 0 else L(g("No. of Equity Shares"))
    sales_ttm, op_ttm = _ttm(q, "Sales"), _ttm(q, "Operating Profit")
    dep_ttm, np_ttm = _ttm(q, "Depreciation"), _ttm(q, "Net profit")
    return {
        "symbol": symbol,
        "file_age_days": (time.time() - os.path.getmtime(path)) / 86400,
        "last_fy": last.date().isoformat(),
        "shares": shares,
        # quality
        "q1": q1, "q2": q2, "q3": q3, "q4": q4, "q5": q5,
        "roa3": roa3, "roe3": roe3, "payout": payout, "neg_roe_sd5": -roe_sd5 if np.isfinite(roe_sd5) else np.nan,
        "fscore": fs, "fscore_n": fn,
        # valuation
        "equity": L(equity), "borrow": L(borrow), "cash": L(cash),
        "sales_ttm": sales_ttm if np.isfinite(sales_ttm) else L(sales),
        "ebitda_ttm": op_ttm if np.isfinite(op_ttm) else L(op),
        "ebit_ttm": (op_ttm - dep_ttm) if np.isfinite(op_ttm) and np.isfinite(dep_ttm) else L(op) - L(dep),
        "np_ttm": np_ttm if np.isfinite(np_ttm) else L(np_),
        "ebitda_5y": _avg_last(op, 5) if op.dropna().size >= 3 else np.nan,
        "roce": L(roce), "roce5": _avg_last(roce, 5) if roce.dropna().size >= 3 else np.nan,
        "cfo": L(cfo), "capex": capex,
    }
